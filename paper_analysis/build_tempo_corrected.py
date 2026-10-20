#!/usr/bin/env python3
"""
build_tempo_corrected.py: pipeline step 5, octave correction.

Turns the raw inputs into Tempo_Corrected.xlsx, the input of tempo_analysis.py.

Inputs
  SOURCE_XLSX  original study workbook, tab "Experimental Data": one row per
               participant with demographics.
               With PRIMARY = "sheet" (default) the BPM columns of this tab
               (Ableton / Tunebat estimates) are the values that get checked.
  ALGO_SUMMARY optional: tempo_summary_*.csv from vocal_tempo.py, used instead
               when PRIMARY = "algorithm".
  TAPS_CSV     export of tempo_tapper.html: one row per recording, file names
               "<row>_<country>_<date>_<song type>.wav", skipped = 1 for clips
               judged too short or unclear.

Rule
  The algorithm gives the tempo; the hand tap decides the metrical level only. For each recording,
  ratio = analysed BPM / tapped BPM is assigned to the nearest of
  1:3, 1:2, 2:3, 1:1, 3:2, 2:1, 3:1 (log scale).
    2:1  -> analysed value is double time -> corrected = analysed / 2
    1:2  -> analysed value is half time   -> corrected = analysed x 2
    3:1, 1:3 -> manual decision required (see MANUAL_DECISIONS)
    other -> unchanged
  Skipped clips are excluded from analysis (no corrected value).

Output (one workbook)
  Tempo table      one row per recording: analysed, tapped, corrected BPM, status
  Tempo summary    corrected tempo by song type and by continent
  Octave check     every corrected or reviewed recording with ratio and decision
  Corrected data   the Experimental Data tab plus three corrected BPM columns
  All clips        ratio and nearest relation for every tapped recording
  Excluded         skipped recordings

All values are computed here and written as plain values, so the result is
identical on every machine. To change a decision, edit MANUAL_DECISIONS or
OVERRIDES below and run again.

    python build_tempo_corrected.py
"""

import os
import re
import sys

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

# =============================================================================
# CONFIG
# =============================================================================
# Paths are relative to the project folder (the folder that contains 02_code), so the
# whole project can be moved or shared without editing them.
PROJECT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DATA = os.path.join(PROJECT, "01_data")
SOURCE_XLSX = os.path.join(DATA, "Study Cross-Continental Tempo of Lullabies - Final Data - Experimental and Repository Data.xlsx")
SOURCE_SHEET = "Experimental Data"
# Source of the tempo values that get octave-checked:
#   "sheet"     = BPM columns of the source workbook (Ableton / Tunebat, used in the paper)
#   "algorithm" = tempo_summary_*.csv from vocal_tempo.py
PRIMARY = "sheet"
ALGO_SUMMARY = os.path.join(DATA, "tempo_summary_generic.csv")   # only used if PRIMARY = "algorithm"
ALGO_BPM_COL = "Reported_BPM"
TAPS_CSV = os.path.join(DATA, "tempo_taps_BR_2026-10-01.csv")
OUTPUT_XLSX = os.path.join(PROJECT, "03_output", "Tempo_Corrected.xlsx")

CLEAR_TOL = 0.08   # within 8% of exact 2:1 / 1:2 = "clear", otherwise "borderline" (label only)

# Decisions for triple relations (3:1 / 1:3), which the rule does not resolve.
# key: file name, value: "halve", "double" or "keep". Reason documented per entry.
MANUAL_DECISIONS = {
    # tap 156.9 vs Ableton 62.16 (ratio 0.40): judged by ear as duple metre, so double
    "0075_India_Jul._12_2021_sleep_song.wav": "double",
}

# Optional: override the automatic rule for single files ("halve", "double", "keep").
OVERRIDES = {}
# =============================================================================

SONGS = {"lullaby": ("Lullaby", "Lullaby BPM"),
         "love_song": ("Love song", "Love Song BPM"),
         "sleep_song": ("Sleep song", "Sleep Song BPM")}
RELATIONS = {"1:3": 1 / 3, "1:2": 1 / 2, "2:3": 2 / 3, "1:1": 1.0, "3:2": 3 / 2, "2:1": 2.0, "3:1": 3.0}
ACTION_FACTOR = {"halve": 0.5, "double": 2.0, "keep": 1.0}

F = lambda **k: Font(name="Arial", **k)
NAVY, BAND, AMBER, GREY = "1F3A5F", "F3F6F9", "FCE9C8", "ECEEF1"
HDR = PatternFill("solid", fgColor="DDE3EA")
LINE = Side(style="thin", color="C8D0D9")


def nearest(ratio):
    k = min(RELATIONS, key=lambda r: abs(np.log(ratio / RELATIONS[r])))
    return k, abs(ratio / RELATIONS[k] - 1)


def load():
    for p in (SOURCE_XLSX, TAPS_CSV) + ((ALGO_SUMMARY,) if PRIMARY == "algorithm" else ()):
        if not os.path.exists(p):
            sys.exit(f"Input not found: {p}")
    src = pd.read_excel(SOURCE_XLSX, sheet_name=SOURCE_SHEET, dtype=str).reset_index(drop=True)
    src.insert(0, "row", src.index + 1)

    taps = pd.read_csv(TAPS_CSV, dtype=str)
    if "pass" in taps.columns:
        taps = taps[pd.to_numeric(taps["pass"], errors="coerce").fillna(1) == 1]
    taps = taps.sort_values("tapped_iso").drop_duplicates("filename", keep="last")
    m = taps["filename"].str.extract(r"^(\d+)_.*_(lullaby|love_song|sleep_song)\.wav$")
    if m.isna().any().any():
        bad = taps.loc[m.isna().any(axis=1), "filename"].tolist()
        sys.exit(f"Unexpected file names in tap CSV: {bad[:5]}")
    taps["row"] = m[0].astype(int)
    taps["song_key"] = m[1]
    taps["tap"] = pd.to_numeric(taps["tapped_bpm"], errors="coerce")
    taps["skipped"] = taps["skipped"].astype(str).isin(["1", "true", "True"])
    taps["spread"] = (pd.to_numeric(taps["sd_ibi_ms"], errors="coerce")
                      / pd.to_numeric(taps["median_ibi_ms"], errors="coerce"))

    if PRIMARY == "sheet":
        s = src.set_index("row")
        taps["analysed_raw"] = [s.at[r, SONGS[k][1]] for r, k in zip(taps["row"], taps["song_key"])]
        taps["analysed"] = pd.to_numeric(taps["analysed_raw"], errors="coerce")
        return src, taps
    algo = pd.read_csv(ALGO_SUMMARY)
    if "Filename" not in algo.columns or ALGO_BPM_COL not in algo.columns:
        sys.exit(f"{ALGO_SUMMARY} needs the columns Filename and {ALGO_BPM_COL}")
    algo = algo.drop_duplicates("Filename").set_index("Filename")[ALGO_BPM_COL]
    missing = sorted(set(taps["filename"]) - set(algo.index))
    if missing:
        print(f"Warning: {len(missing)} tapped files have no algorithm estimate, e.g. {missing[:3]}")
    taps["analysed_raw"] = taps["filename"].map(algo)
    taps["analysed"] = pd.to_numeric(taps["analysed_raw"], errors="coerce")
    return src, taps


def decide(t):
    rows = []
    for _, r in t.iterrows():
        d = {"ratio": np.nan, "nearest": "", "deviation": np.nan, "tier": "", "action": "keep",
             "reason": "", "status": "Included"}
        if r.skipped:
            d.update(status="Excluded", reason="skipped during tapping")
        elif pd.isna(r.analysed) or pd.isna(r.tap) or r.tap <= 0:
            d.update(status="Excluded", reason="no numeric analysed or tapped value")
        else:
            ratio = r.analysed / r.tap
            rel, dev = nearest(ratio)
            d.update(ratio=ratio, nearest=rel, deviation=dev, tier="clear" if dev <= CLEAR_TOL else "borderline")
            if rel == "2:1":
                d.update(action="halve", reason="rule: nearest 2:1")
            elif rel == "1:2":
                d.update(action="double", reason="rule: nearest 1:2")
            elif rel in ("3:1", "1:3"):
                if r.filename in MANUAL_DECISIONS:
                    d.update(action=MANUAL_DECISIONS[r.filename], reason=f"manual: triple relation {rel}")
                else:
                    d.update(action="keep", reason=f"triple relation {rel}, no manual decision: kept")
            if r.filename in OVERRIDES:
                d.update(action=OVERRIDES[r.filename], reason="override")
        rows.append(d)
    t = pd.concat([t.reset_index(drop=True), pd.DataFrame(rows)], axis=1)
    t["corrected"] = np.where(t["status"] == "Included", t["analysed"] * t["action"].map(ACTION_FACTOR), np.nan)
    t["correction"] = t["action"].map({"halve": "halved", "double": "doubled", "keep": ""}).where(
        t["status"] == "Included", "")
    t["song_type"] = t["song_key"].map(lambda k: SONGS[k][0])
    order = {"lullaby": 0, "love_song": 1, "sleep_song": 2}
    return t.assign(_o=t["song_key"].map(order)).sort_values(["row", "_o"]).drop(columns="_o").reset_index(drop=True)


def header(ws, row, labels, fill=HDR, white=False):
    for j, h in enumerate(labels, 1):
        c = ws.cell(row, j, h)
        c.font = F(bold=True, color="FFFFFF" if white else "000000")
        c.fill = fill
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def widths(ws, ws_widths):
    for j, w in enumerate(ws_widths, 1):
        ws.column_dimensions[get_column_letter(j)].width = w


def val(x):
    if x is None or (isinstance(x, float) and np.isnan(x)):
        return None
    if isinstance(x, (np.floating,)):
        return float(x)
    if isinstance(x, (np.integer,)):
        return int(x)
    return x


def describe(x):
    x = np.asarray(x, float)
    x = x[~np.isnan(x)]
    return dict(n=len(x), mean=x.mean() if len(x) else np.nan, sd=x.std(ddof=1) if len(x) > 1 else np.nan,
                median=np.median(x) if len(x) else np.nan, min=x.min() if len(x) else np.nan,
                max=x.max() if len(x) else np.nan)


def build(src, t):
    wb = Workbook()
    ed = src.set_index("row")

    # ---- Tempo table ----
    tt = wb.active
    tt.title = "Tempo table"
    tt["A1"] = "Tempo per recording: analysed, hand-tapped and octave-corrected BPM"
    tt["A1"].font = F(bold=True, size=13, color=NAVY)
    tt["A2"] = ("Corrected BPM = analysed BPM, halved or doubled where its ratio to the hand tap was closer to "
                "2:1 or 1:2 than to 3:2, 2:3 or 1:1. Highlighted rows were corrected; grey rows are excluded.")
    tt["A2"].font = F(italic=True, size=9, color="555555")
    H = ["ID", "Continent", "Country", "Song type", "Analysed BPM", "Tapped BPM", "Corrected BPM",
         "Correction", "Status", "Gender", "Age", "Musical training"]
    header(tt, 4, H, PatternFill("solid", fgColor=NAVY), white=True)
    tt.row_dimensions[4].height = 30
    for i, r in enumerate(t.itertuples(), 5):
        p = ed.loc[r.row]
        an = r.analysed if not pd.isna(r.analysed) else r.analysed_raw
        vals = [r.row, p["Continent"], p["Country of origin"], r.song_type, an,
                None if r.skipped else r.tap, r.corrected, r.correction, r.status,
                p["What is your gender?"], p["How old are you?"], p["Have you had any formal musical training?"]]
        for j, v in enumerate(vals, 1):
            c = tt.cell(i, j, val(v))
            c.font = F(size=10, bold=(j == 7))
            c.border = Border(bottom=LINE)
            c.alignment = Alignment(horizontal="center" if j in (1, 4, 5, 6, 7, 8, 9, 11) else "left")
            if j in (5, 6, 7):
                c.number_format = "0.0"
    end = 4 + len(t)
    tt.conditional_formatting.add(f"A5:L{end}", FormulaRule(formula=['$I5="Excluded"'],
                                  font=Font(name="Arial", color="9AA3AD", italic=True),
                                  fill=PatternFill("solid", fgColor=GREY), stopIfTrue=True))
    tt.conditional_formatting.add(f"A5:L{end}", FormulaRule(formula=['$H5<>""'], fill=PatternFill("solid", fgColor=AMBER)))
    tt.conditional_formatting.add(f"A5:L{end}", FormulaRule(formula=["MOD($A5,2)=0"], fill=PatternFill("solid", fgColor=BAND)))
    tt.freeze_panes = "A5"
    tt.auto_filter.ref = f"A4:L{end}"
    tt.sheet_view.showGridLines = False
    widths(tt, [6, 15, 18, 11, 12, 11, 13, 11, 11, 10, 9, 24])

    # ---- Tempo summary ----
    sm = wb.create_sheet("Tempo summary")
    sm["A1"] = "Corrected tempo by song type and continent (included recordings)"
    sm["A1"].font = F(bold=True, size=13, color=NAVY)
    sm.sheet_view.showGridLines = False
    inc = t[t["status"] == "Included"].copy()
    inc["continent"] = inc["row"].map(ed["Continent"].str.strip())
    inc["country"] = inc["row"].map(ed["Country of origin"])
    cols = ["Group", "n", "Mean BPM", "SD", "Median BPM", "Min", "Max", "Corrected"]
    rw = 3

    def block(title, groups):
        nonlocal rw
        sm.cell(rw, 1, title).font = F(bold=True, size=11, color=NAVY)
        rw += 1
        header(sm, rw, cols, PatternFill("solid", fgColor=NAVY), white=True)
        rw += 1
        for name, g in groups:
            d = describe(g["corrected"])
            for j, v in enumerate([name, d["n"], d["mean"], d["sd"], d["median"], d["min"], d["max"],
                                   int((g["correction"] != "").sum())], 1):
                c = sm.cell(rw, j, val(v))
                c.font = F(size=10)
                c.border = Border(bottom=LINE)
                if j in (3, 4, 5, 6, 7):
                    c.number_format = "0.0"
            rw += 1
        rw += 1

    block("By song type", [(s[0], inc[inc.song_type == s[0]]) for s in SONGS.values()])
    conts = [c for c in ["Africa", "Asia", "Europe", "North America", "South America"] if c in set(inc.continent)]
    for s in SONGS.values():
        sub = inc[inc.song_type == s[0]]
        groups = [(c, sub[(sub.continent == c) & (sub.country != "Australia")]) for c in conts]
        if (sub.country == "Australia").any():
            groups.append(("Australia*", sub[sub.country == "Australia"]))
        block(f"{s[0]} by continent", groups)
    sm.cell(rw, 1, "*Australia is labelled 'Asia ' in the source sheet and listed separately here.").font = F(
        italic=True, size=9, color="555555")
    widths(sm, [24, 8, 12, 10, 12, 10, 10, 11])

    # ---- Octave check ----
    oc = wb.create_sheet("Octave check")
    oc["A1"] = ("Recordings whose analysed BPM was corrected or needed review. ratio = analysed / tapped; "
                "nearest = closest of 1:3, 1:2, 2:3, 1:1, 3:2, 2:1, 3:1.")
    oc["A1"].font = F(italic=True)
    hd = ["filename", "row", "song_type", "tapped_bpm", "analysed_bpm", "ratio", "nearest", "deviation",
          "tier", "action", "reason", "corrected_bpm", "tap_spread"]
    header(oc, 3, hd)
    sel = t[(t["status"] == "Included") & (t["nearest"].isin(["2:1", "1:2", "3:1", "1:3"]) | (t["reason"] == "override"))]
    for i, r in enumerate(sel.sort_values(["tier", "deviation"]).itertuples(), 4):
        for j, v in enumerate([r.filename, r.row, r.song_type, r.tap, r.analysed, r.ratio, r.nearest,
                               r.deviation, r.tier, r.action, r.reason, r.corrected, r.spread], 1):
            c = oc.cell(i, j, val(v))
            c.font = F()
            if j in (6, 13):
                c.number_format = "0.000"
            if j == 8:
                c.number_format = "0.0%"
            if j == 12:
                c.number_format = "0.00"
    oc.freeze_panes = "A4"
    widths(oc, [44, 6, 11, 11, 13, 9, 9, 10, 11, 9, 30, 13, 11])

    # ---- Corrected data ----
    cd = wb.create_sheet("Corrected data")
    base_cols = list(src.columns)
    new_cols = ["Lullaby BPM corrected", "Love Song BPM corrected", "Sleep Song BPM corrected", "Octave corrections in row"]
    header(cd, 1, base_cols)
    for j, h in enumerate(new_cols, len(base_cols) + 1):
        c = cd.cell(1, j, h)
        c.font = F(bold=True)
        c.fill = PatternFill("solid", fgColor="C9E7D3")
    lookup = t.set_index(["row", "song_key"])
    for i, (_, r) in enumerate(src.iterrows(), 2):
        for j, h in enumerate(base_cols, 1):
            v = r[h]
            if h in [s[1] for s in SONGS.values()]:
                num = pd.to_numeric(v, errors="coerce")
                v = num if pd.notna(num) else v
            if h == "row":
                v = int(v)
            cd.cell(i, j, val(None if (isinstance(v, float) and np.isnan(v)) else v)).font = F()
        n_corr = 0
        for k, (key, (_, col)) in enumerate(SONGS.items()):
            orig = pd.to_numeric(r[col], errors="coerce")
            if (r["row"], key) in lookup.index:
                rec = lookup.loc[(r["row"], key)]
                out = "excluded" if rec["status"] == "Excluded" else rec["corrected"]
                n_corr += int(rec["correction"] != "")
            else:  # no recording: mirror the source cell (number, "X" or empty)
                raw = r[col]
                out = orig if pd.notna(orig) else (raw if isinstance(raw, str) and raw.strip() else None)
            c = cd.cell(i, len(base_cols) + 1 + k, val(out))
            c.font = F()
            c.number_format = "0.00"
        cd.cell(i, len(base_cols) + 4, n_corr).font = F()
    cd.freeze_panes = "B2"
    widths(cd, [16] * (len(base_cols) + 4))

    # ---- All clips ----
    ac = wb.create_sheet("All clips")
    header(ac, 1, ["filename", "song_type", "status", "tapped_bpm", "analysed_bpm", "ratio", "nearest",
                   "deviation", "action", "corrected_bpm", "tap_spread"])
    for i, r in enumerate(t.itertuples(), 2):
        for j, v in enumerate([r.filename, r.song_type, r.status, None if r.skipped else r.tap, r.analysed,
                               r.ratio, r.nearest, r.deviation, r.action if r.status == "Included" else "",
                               r.corrected, r.spread], 1):
            c = ac.cell(i, j, val(v))
            c.font = F()
            if j in (6, 11):
                c.number_format = "0.000"
            if j == 8:
                c.number_format = "0.0%"
    ac.freeze_panes = "A2"
    ac.auto_filter.ref = f"A1:K{len(t) + 1}"
    widths(ac, [44, 11, 10, 11, 13, 9, 9, 10, 9, 13, 11])

    # ---- Excluded ----
    ex = wb.create_sheet("Excluded")
    header(ex, 1, ["filename", "song_type", "reason", "analysed_raw"])
    for i, r in enumerate(t[t["status"] == "Excluded"].itertuples(), 2):
        for j, v in enumerate([r.filename, r.song_type, r.reason, r.analysed_raw], 1):
            ex.cell(i, j, val(v)).font = F()
    widths(ex, [44, 11, 40, 14])

    os.makedirs(os.path.dirname(OUTPUT_XLSX), exist_ok=True)
    wb.save(OUTPUT_XLSX)


def main():
    src, taps = load()
    t = decide(taps)
    build(src, t)
    inc = t[t["status"] == "Included"]
    print(f"Recordings: {len(t)}  included: {len(inc)}  excluded: {int((t['status'] == 'Excluded').sum())}")
    print(f"Corrections: {int((inc['correction'] == 'halved').sum())} halved, "
          f"{int((inc['correction'] == 'doubled').sum())} doubled")
    review = inc[inc["nearest"].isin(["3:1", "1:3"])]
    for _, r in review.iterrows():
        print(f"Triple relation: {r.filename} -> {r.action} ({r.reason})")
    print(f"Saved {OUTPUT_XLSX}")


if __name__ == "__main__":
    main()
